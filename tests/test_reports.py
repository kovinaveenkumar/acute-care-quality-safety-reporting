import zipfile

import openpyxl

from qsr import build_reports
from qsr.config import OUT_DIR


def test_reports_build_and_contain_key_content(pipeline):
    build_reports.main()
    html = (OUT_DIR / "dashboard.html").read_text()
    assert "Executive scorecard" in html and "Prepared by" in html and "<svg" in html
    assert "MSW" in html and "Control charts" in html
    wb = openpyxl.load_workbook(OUT_DIR / "Quality_Committee_Pack.xlsx")
    for name in ("Scorecard", "Findings", "SPC Data", "Control Charts", "Abstraction Tracker", "Public Reporting", "Data Quality Log"):
        assert name in wb.sheetnames
    assert wb["Scorecard"].max_row > 20
    assert (OUT_DIR / "Quality_Brief.pdf").read_bytes()[:4] == b"%PDF"
    with zipfile.ZipFile(OUT_DIR / "Quality_Committee_Pack.xlsx") as z:
        assert any(n.startswith("xl/charts/") for n in z.namelist())
